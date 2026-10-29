# Resolve Google Maps short links from the GEO Excel into lat/lng.
from __future__ import annotations

import json
import re
import ssl
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / "scripts" / "geo-cache.json"
XLSX = ROOT / "PRESTADORES TURÍSTICOS HABILITADOS con GEO.xlsx"

UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
)

# Place pin - !3dLAT!4dLNG is the dropped-pin, not the camera center.
PIN_RE = re.compile(r"!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)")
# Place path - /place/-42.1,-71.2 or /search/-42.1,-71.2
PLACE_RE = re.compile(r"/(?:place|search)/(-?\d+\.\d+)\s*,\s*(-?\d+\.\d+)")
# Query - q=-42.1,-71.2 or query=-42.1,-71.2
QUERY_RE = re.compile(r"[?&](?:q|query)=(-?\d+\.\d+)\s*,\s*(-?\d+\.\d+)")
# Camera - /@-42.1,-71.2,17z last resort
AT_RE = re.compile(r"@(-?\d+\.\d+),(-?\d+\.\d+)")


def load_cache() -> dict:
    if CACHE.exists():
        return json.loads(CACHE.read_text(encoding="utf-8"))
    return {}


def save_cache(cache: dict) -> None:
    CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def parse_coords(url: str) -> tuple[float, float] | None:
    text = urllib.parse.unquote_plus(url.replace("%21", "!"))
    for pattern in (PIN_RE, PLACE_RE, QUERY_RE, AT_RE):
        match = pattern.search(text)
        if match:
            lat, lng = float(match.group(1)), float(match.group(2))
            if -50 < lat < -40 and -80 < lng < -60:
                return round(lat, 6), round(lng, 6)
    return None


def fetch_final(url: str) -> str:
    ctx = ssl.create_default_context()
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept-Language": "es-AR,es;q=0.9"})
    try:
        with urllib.request.urlopen(req, timeout=25, context=ctx) as resp:
            final = resp.geturl()
            body = resp.read(80000).decode("utf-8", errors="ignore")
            return final + "\n" + body
    except urllib.error.HTTPError as err:
        loc = err.headers.get("Location") or ""
        return (loc or "") + "\n" + (err.read(40000).decode("utf-8", errors="ignore") if err.fp else "")


def resolve(url: str) -> dict:
    blob = fetch_final(url)
    coords = parse_coords(blob)
    # Second hop if the first response is still a short/app link.
    if not coords:
        hrefs = re.findall(r"https?://(?:www\.)?(?:google\.[^/\"'\s]+/maps|maps\.google\.[^/\"'\s]+)[^\"'\s<>]+", blob)
        for href in hrefs[:6]:
            coords = parse_coords(href)
            if coords:
                blob = href
                break
            hop = fetch_final(href)
            coords = parse_coords(hop)
            if coords:
                blob = hop
                break
    if not coords:
        return {"ok": False, "url": url, "sample": blob[:400]}
    return {"ok": True, "url": url, "lat": coords[0], "lng": coords[1], "resolved": blob.split("\n", 1)[0][:300]}


def main() -> None:
    from openpyxl import load_workbook

    wb = load_workbook(XLSX, data_only=True)
    urls = []
    for sheet in ("Alojamientos", "Gastronomía", "Otros Servicos"):
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, str) and "maps.app.goo.gl" in cell:
                    urls.append(cell.strip())
    unique = list(dict.fromkeys(urls))
    print("unique urls", len(unique))
    cache = load_cache()
    for i, url in enumerate(unique, 1):
        if url in cache and cache[url].get("ok"):
            print(i, "cached", url, cache[url]["lat"], cache[url]["lng"])
            continue
        try:
            result = resolve(url)
        except Exception as exc:
            result = {"ok": False, "url": url, "error": str(exc)}
        cache[url] = result
        save_cache(cache)
        print(i, "ok" if result.get("ok") else "FAIL", url, result.get("lat"), result.get("lng"), result.get("error") or result.get("sample", "")[:80])
        time.sleep(0.4)
    ok = sum(1 for v in cache.values() if v.get("ok"))
    print("resolved", ok, "of", len(unique))


if __name__ == "__main__":
    main()
