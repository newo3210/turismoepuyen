# Drive import - compress maps/artesanos/stupa photos and rebuild listings from the Excel.
from __future__ import annotations

import hashlib
import json
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent))
from apply_directorio import normalize_phone

ROOT = Path(__file__).resolve().parents[1]
MIRROR = ROOT / "mirror"
DRIVE = ROOT / "googledrive" / "PÁGINA WEB"
GEO_XLSX = ROOT / "PRESTADORES TURÍSTICOS HABILITADOS con GEO.xlsx"
GEO_CACHE = ROOT / "scripts" / "geo-cache.json"

# Puerto Bonito services - the Excel has no GEO column there; reuse the park pin.
SERVICE_GEO = {
    "puerto-bonito": "https://maps.app.goo.gl/nSPei7dt8uNuMfoM6",
    "puerto-bonito-aventura": "https://maps.app.goo.gl/nSPei7dt8uNuMfoM6",
}

PAGES = {
    "el-esquilador": "/el-esquilador/",
    "piuke": "/cabana-piuke/",
    "el-refugio-del-lago": "/refugio-del-lago/",
    "cabanas-del-epuyen": "/cabanas-del-epuyen/",
    "runan-nativo": "/runan-nativo/",
    "carelhue": "/cabanas-carelhue/",
    "cabanas-carelhue": "/cabanas-carelhue/",
}

IMAGES = {
    "Cabañas": "/wp-content/uploads/2024/04/alojamientos.jpeg",
    "Camping": "/wp-content/uploads/2024/05/lago-epuyen.jpg",
    "Hostería": "/wp-content/uploads/2024/04/Hosteria-foto-principal.webp",
    "Casa": "/wp-content/uploads/2024/04/conocenos.webp",
    "Hostel": "/wp-content/uploads/2024/04/actividades-1.jpeg",
    "Dormis": "/wp-content/uploads/2024/12/nogales-interior-768x512.webp",
    "Complejo": "/wp-content/uploads/2024/04/alojamientos.jpeg",
    "Kayak": "/wp-content/uploads/2024/05/lago-epuyen.jpg",
    "Guía": "/wp-content/uploads/2024/04/actividades-1.jpeg",
    "Agencia": "/wp-content/uploads/2024/04/actividades-1.jpeg",
    "Remís": "/wp-content/uploads/2024/04/maapa-cabezal.jpeg",
    "Restaurante": "/wp-content/uploads/2024/04/Hosteria-foto-principal.webp",
    "Confitería": "/wp-content/uploads/2024/04/conocenos.webp",
    "Rotisería": "/wp-content/uploads/2024/04/conocenos.webp",
    "Pizzería": "/wp-content/uploads/2024/04/conocenos.webp",
    "Cafetería": "/wp-content/uploads/2024/04/conocenos.webp",
    "Cervecería": "/wp-content/uploads/2024/04/conocenos.webp",
    "Fábrica": "/wp-content/uploads/2024/04/conocenos.webp",
    "Chocolatería": "/wp-content/uploads/2024/04/conocenos.webp",
    "Regional": "/wp-content/uploads/2024/04/conocenos.webp",
    "Restobar": "/wp-content/uploads/2024/04/conocenos.webp",
    "Heladería": "/wp-content/uploads/2024/04/conocenos.webp",
    "Pastelería": "/wp-content/uploads/2024/04/conocenos.webp",
    "Vinoteca": "/wp-content/uploads/2024/04/conocenos.webp",
}

ZONES = {
    "zona lago": (-42.217, -71.392),
    "zona el lago": (-42.217, -71.392),
    "lago": (-42.217, -71.392),
    "la rinconada": (-42.208, -71.412),
    "riconada": (-42.208, -71.412),
    "casco urbano (céntrico)": (-42.2305, -71.3575),
    "casco urbano (centrico)": (-42.2305, -71.3575),
    "cascco urbano (céntrico )": (-42.2305, -71.3575),
    "casco urbano": (-42.2305, -71.3575),
    "casco urbano (barrio alto)": (-42.233, -71.350),
    "el pedregoso": (-42.176, -71.355),
    "zona vialidad": (-42.246, -71.348),
    "casco urbano (ruta 40)": (-42.236, -71.352),
    "ruta 40": (-42.228, -71.340),
    "zona puente de los vecinos": (-42.205, -71.372),
    "caco urbano (céntrico)": (-42.2305, -71.3575),
    "caco urbano (centrico)": (-42.2305, -71.3575),
}
DEFAULT_LL = (-42.230, -71.365)

CAT_MAP = {
    "cabañas": "Cabañas",
    "cabaña": "Cabañas",
    "camping": "Camping",
    "hosteria": "Hostería",
    "hostería": "Hostería",
    "restorant": "Hostería",
    "vivienda turistica": "Casa",
    "vivienda turística": "Casa",
    "hostel": "Hostel",
    "dormis": "Dormis",
    "complejo": "Complejo",
    "alquiler de kayak": "Kayak",
    "guia de turismo residente": "Guía",
    "guía de turismo residente": "Guía",
    "agencias de viajes virtual": "Agencia",
    "remis": "Remís",
    "remís": "Remís",
    "restaurante": "Restaurante",
    "restaurant": "Restaurante",
    "confiteria": "Confitería",
    "confitería": "Confitería",
    "rotiseria": "Rotisería",
    "rotisería": "Rotisería",
    "pizzeria": "Pizzería",
    "pizzería": "Pizzería",
    "heladeria": "Heladería",
    "heladería": "Heladería",
    "vinoteca": "Vinoteca",
    "fabrica de cerveza": "Cervecería",
    "fábrica de cerveza": "Cervecería",
    "cerveceria": "Cervecería",
    "cervecería": "Cervecería",
    "fabrica artesanal": "Fábrica",
    "fábrica artesanal": "Fábrica",
    "fabrica de alfajores": "Fábrica",
    "fábrica de alfajores": "Fábrica",
    "cafeteria-pasteleria": "Cafetería",
    "cafeteria": "Cafetería",
    "cafetería": "Cafetería",
    "pasteleria": "Pastelería",
    "pastelería": "Pastelería",
    "chocolateria": "Chocolatería",
    "chocolatería": "Chocolatería",
    "fabrica de licores": "Fábrica",
    "artículos regionales": "Regional",
    "articulos regionales": "Regional",
    "restobar": "Restobar",
}


# Excel cell - flatten numbers like 2945469840.0 and empty cells.
def cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


# Display names - title case with small Spanish particles left lowercase.
def titlecase(name: str) -> str:
    small = {"de", "del", "la", "el", "los", "las", "y", "da"}
    words = name.strip().split()
    out = []
    for i, word in enumerate(words):
        low = word.lower()
        if i and low in small:
            out.append(low)
        else:
            out.append(word[:1].upper() + word[1:].lower() if word else word)
    return " ".join(out)


def slugify(name: str) -> str:
    text = name.lower()
    repl = {"á": "a", "é": "e", "í": "i", "ó": "o", "ú": "u", "ñ": "n", "ü": "u"}
    text = "".join(repl.get(ch, ch) for ch in text)
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text or "item"


def cats_from_rubro(raw: str) -> list[str]:
    parts = re.split(r"[/,]", raw or "")
    out = []
    for part in parts:
        key = part.strip().lower()
        if not key:
            continue
        label = CAT_MAP.get(key, titlecase(part.strip()))
        if label not in out:
            out.append(label)
    return out


def zone_ll(zone: str, key: str) -> tuple[float, float]:
    zkey = re.sub(r"\s+", " ", zone.lower().strip())
    lat, lon = ZONES.get(zkey, DEFAULT_LL)
    digest = int(hashlib.md5(key.encode()).hexdigest()[:6], 16)
    dlat = ((digest % 80) - 40) * 0.00008
    dlon = (((digest // 80) % 80) - 40) * 0.00008
    return round(lat + dlat, 6), round(lon + dlon, 6)


# Geo cache - lat/lng resolved from the Excel maps.app.goo.gl links.
def load_geo_cache() -> dict:
    if GEO_CACHE.exists():
        return json.loads(GEO_CACHE.read_text(encoding="utf-8"))
    return {}


# Pin coords - prefer the Excel Google Maps link; otherwise jitter the zone centroid.
def coords_from_geo(geo_url: str, zone: str, key: str, cache: dict) -> tuple[float, float, bool, str]:
    geo_url = (geo_url or "").strip()
    hit = cache.get(geo_url) if geo_url else None
    if hit and hit.get("ok"):
        return float(hit["lat"]), float(hit["lng"]), False, geo_url
    lat, lng = zone_ll(zone or "", key)
    return lat, lng, True, geo_url


def phones_from(raw: str) -> list[dict]:
    out = []
    seen = set()
    for part in re.split(r"[/]", raw or ""):
        norm = normalize_phone(part)
        if not norm or norm["wa"] in seen:
            continue
        seen.add(norm["wa"])
        out.append(norm)
    return out


# Web JPEG - shrink Drive originals so the static site can ship them.
def save_web(src: Path, dest: Path, max_side: int = 1800, quality: int = 82) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    image = Image.open(src)
    if image.mode in ("RGBA", "P"):
        image = image.convert("RGB")
    else:
        image = image.convert("RGB")
    image.thumbnail((max_side, max_side))
    image.save(dest, "JPEG", quality=quality, optimize=True)
    print("img", dest.relative_to(MIRROR), dest.stat().st_size // 1024, "KB")


def copy_pdf(src: Path, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)
    print("pdf", dest.relative_to(MIRROR))


def export_media() -> None:
    artesanos = DRIVE / "32° Encuentro Provincial de Artesanos"
    mapas = DRIVE / "MAPAS"
    stupa = DRIVE / "STUPA SAMANTABHADRA"
    save_web(mapas / "MAPA EPUYÉN.png", MIRROR / "media" / "mapas" / "mapa-epuyen.jpg", 2400, 85)
    save_web(mapas / "FOLLETO TURISMO INVIERNO.png", MIRROR / "media" / "mapas" / "folleto-invierno.jpg", 2400, 85)
    save_web(artesanos / "CHRISTIAN HERRERA.jpg", MIRROR / "media" / "artesanos" / "christian-herrera.jpg")
    rancheros = next(artesanos.glob("*RANCHEROS*"))
    save_web(rancheros, MIRROR / "media" / "artesanos" / "super-rancheros.jpg")
    save_web(artesanos / "PREMIO MEJOR OBRA ARTESANAL.jpg", MIRROR / "media" / "artesanos" / "premio-mejor-obra.jpg")
    save_web(artesanos / "DJI_20260404233714_0105_D.jpg", MIRROR / "media" / "artesanos" / "escenario-noche.jpg")
    copy_pdf(artesanos / "RESUMEN ENCUENTRO DE ARTESANOS.pdf", MIRROR / "media" / "artesanos" / "resumen-32.pdf")
    save_web(stupa / "STUPA.JPG", MIRROR / "media" / "stupa" / "stupa.jpg", 1600, 82)
    copy_pdf(stupa / "STUPA DE EPUYÉN - RESÚMEN.pdf", MIRROR / "media" / "stupa" / "resumen.pdf")


def merge_rows(rows: list[dict], kind: str, cache: dict) -> list[dict]:
    by: dict[str, dict] = {}
    fallback = IMAGES["Cabañas"] if kind == "stay" else IMAGES["Guía"]
    if kind == "food":
        fallback = IMAGES["Restaurante"]
    for row in rows:
        name = titlecase(row["name"] or row["responsable"])
        if not name:
            continue
        sid = slugify(row["name"] or row["responsable"])
        item = by.get(sid)
        cats = cats_from_rubro(row["rubro"])
        phones = phones_from(row["contact"])
        mail = row["mail"]
        notes = ""
        if mail and "@" not in mail:
            notes = mail
            mail = ""
        geo_url = (row.get("geo") or "").strip() or SERVICE_GEO.get(sid, "")
        if not item:
            lat, lng, approx, maps_url = coords_from_geo(geo_url, row.get("zone") or "", sid, cache)
            item = {
                "id": sid,
                "name": name,
                "categories": [],
                "responsable": titlecase(row["responsable"]) if row.get("responsable") else "",
                "address": titlecase(row["address"]) if row.get("address") else "",
                "zone": (row.get("zone") or "").strip(),
                "email": mail,
                "notes": notes,
                "phones": phones,
                "lat": lat,
                "lng": lng,
                "page": PAGES.get(sid, ""),
                "image": IMAGES.get(cats[0] if cats else "", fallback),
                "coordsApproximate": approx,
            }
            if maps_url:
                item["mapsUrl"] = maps_url
            by[sid] = item
        elif geo_url and item.get("coordsApproximate"):
            lat, lng, approx, maps_url = coords_from_geo(geo_url, item.get("zone") or "", sid, cache)
            item["lat"] = lat
            item["lng"] = lng
            item["coordsApproximate"] = approx
            if maps_url:
                item["mapsUrl"] = maps_url
        for cat in cats:
            if cat not in item["categories"]:
                item["categories"].append(cat)
        have = {p["wa"] for p in item["phones"]}
        for phone in phones:
            if phone["wa"] not in have:
                item["phones"].append(phone)
        if mail and not item["email"]:
            item["email"] = mail
        if notes and not item["notes"]:
            item["notes"] = notes
        if row.get("address") and not item["address"]:
            item["address"] = titlecase(row["address"])
        if row.get("zone") and not item["zone"]:
            item["zone"] = row["zone"].strip()
    items = list(by.values())
    for item in items:
        item["whatsapp"] = item["phones"][0]["wa"] if item["phones"] else ""
        if not item["categories"]:
            item["categories"] = ["Servicio"] if kind == "service" else ["Cabañas"]
        item["image"] = IMAGES.get(item["categories"][0], item["image"])
        if not item.get("notes"):
            item.pop("notes", None)
        if not item.get("phones") and not item.get("email"):
            item["pendingDetails"] = True
        if item.get("coordsApproximate"):
            item.pop("mapsUrl", None)
    items.sort(key=lambda row: row["name"])
    return items


def export_json() -> None:
    xlsx = GEO_XLSX if GEO_XLSX.exists() else DRIVE / "PRESTADORES TURÍSTICOS HABILITADOS.xlsx"
    cache = load_geo_cache()
    wb = load_workbook(xlsx, data_only=True)
    stay_rows = []
    ws = wb["Alojamientos"]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        vals = [cell(c) for c in (list(row[:7]) + [""] * 7)[:7]]
        name, rubro, contact, mail, address, zone, geo = vals
        if i < 3 or not name or name.upper().startswith("NOMBRE"):
            continue
        stay_rows.append(
            {
                "name": name,
                "rubro": rubro,
                "contact": contact,
                "mail": mail,
                "address": address,
                "zone": zone,
                "geo": geo,
                "responsable": "",
            }
        )
    food_rows = []
    ws = wb["Gastronomía"]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        vals = [cell(c) for c in (list(row[:7]) + [""] * 7)[:7]]
        name, rubro, contact, mail, address, zone, geo = vals
        if i < 3 or not name or name.upper().startswith("NOMBRE"):
            continue
        food_rows.append(
            {
                "name": name,
                "rubro": rubro,
                "contact": contact,
                "mail": mail,
                "address": address,
                "zone": zone,
                "geo": geo,
                "responsable": "",
            }
        )
    named_guides = {
        "paoepuyen@gmail.com": "Rossi Paola Andrea",
        "a.verdasco89@gmail.com": "Verdasco Antonela",
    }
    svc_rows = []
    ws = wb["Otros Servicos"]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        vals = [cell(c) for c in (list(row[:4]) + [""] * 4)[:4]]
        name, contact, mail, rubro = vals
        if i < 3 or (not name and not contact):
            continue
        if "AUTO" in mail.upper():
            name = name or "Chofer Pablo"
        if mail.lower() in named_guides:
            name = named_guides[mail.lower()]
        if not name:
            name = "Guía"
        svc_rows.append(
            {
                "name": name,
                "rubro": rubro,
                "contact": contact,
                "mail": mail,
                "address": "",
                "zone": "",
                "geo": "",
                "responsable": name,
            }
        )
    stay = merge_rows(stay_rows, "stay", cache)
    food = merge_rows(food_rows, "food", cache)
    svc = merge_rows(svc_rows, "service", cache)
    (MIRROR / "data").mkdir(exist_ok=True)
    (MIRROR / "data" / "alojamientos.json").write_text(
        json.dumps({"kind": "stay", "items": stay}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    (MIRROR / "data" / "gastronomia.json").write_text(
        json.dumps({"kind": "food", "items": food}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    (MIRROR / "data" / "servicios.json").write_text(
        json.dumps({"kind": "service", "items": svc}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    exact = lambda items: sum(1 for item in items if not item.get("coordsApproximate"))
    print("xlsx", xlsx.name)
    print("stay", len(stay), "exact", exact(stay))
    print("food", len(food), "exact", exact(food))
    print("svc", len(svc), "exact", exact(svc))


def main() -> None:
    if "--json" not in sys.argv:
        export_media()
    export_json()


if __name__ == "__main__":
    main()
